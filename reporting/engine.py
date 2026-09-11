"""Tiny document model rendered to DOCX (Word) and to a colourful HTML page printed to PDF.

Content is written once as a list of blocks; both renderers read the same
blocks, so the Word report and the PDF always say the same thing.

Inline markup inside text: **bold**, *italic*, `code`.
"""
from __future__ import annotations

import base64
import html
import re
import subprocess
from dataclasses import dataclass, field
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

PALETTE = {
    "primary": "4F46E5", "violet": "7C3AED", "teal": "0D9488", "amber": "D97706", "red": "DC2626",
    "green": "16A34A", "sky": "0284C7", "ink": "1E293B", "muted": "64748B", "line": "E2E8F0",
    "tint": "EEF2FF", "violet_tint": "F5F3FF",
}
CALLOUT = {  # kind -> (accent, background, default title)
    "key": ("7C3AED", "F5F3FF", "Key insight"),
    "proof": ("0D9488", "F0FDFA", "How we prove it"),
    "info": ("0284C7", "F0F9FF", "Note"),
    "warn": ("D97706", "FFFBEB", "Limitation"),
    "success": ("16A34A", "F0FDF4", "Result"),
}


# ------------------------------------------------------------------- blocks
@dataclass
class Doc:
    title: str
    subtitle: str = ""
    meta: dict = field(default_factory=dict)   # cover fields: team, guide, institution, date, domain
    blocks: list = field(default_factory=list)
    landscape: bool = False

    def h1(self, text): self.blocks.append(("h1", text)); return self
    def h2(self, text): self.blocks.append(("h2", text)); return self
    def h3(self, text): self.blocks.append(("h3", text)); return self
    def p(self, text): self.blocks.append(("p", text)); return self
    def bullets(self, items): self.blocks.append(("bullets", list(items))); return self
    def numbers(self, items): self.blocks.append(("numbers", list(items))); return self
    def code(self, text): self.blocks.append(("code", text)); return self
    def pagebreak(self): self.blocks.append(("pagebreak", None)); return self
    def toc(self): self.blocks.append(("toc", None)); return self

    def callout(self, text, kind="key", title=None):
        self.blocks.append(("callout", {"text": text, "kind": kind, "title": title or CALLOUT[kind][2]}))
        return self

    def kpis(self, items):
        """items: [(value, label, sub)]"""
        self.blocks.append(("kpis", list(items)))
        return self

    def table(self, columns, rows, caption=None, widths=None, highlight=None, font=None, first_col_bold=False):
        self.blocks.append(("table", {"columns": columns, "rows": rows, "caption": caption, "widths": widths,
                                      "highlight": set(highlight or []), "font": font,
                                      "first_col_bold": first_col_bold}))
        return self

    def figure(self, path, caption, width=16.0):
        self.blocks.append(("figure", {"path": Path(path), "caption": caption, "width": width}))
        return self

    def flow(self, steps, caption=None):
        """A left-to-right process diagram: steps = [(title, detail)]"""
        self.blocks.append(("flow", {"steps": steps, "caption": caption}))
        return self


INLINE = re.compile(r"(\*\*.+?\*\*|\*[^*]+?\*|`[^`]+?`)")


def _runs(text: str):
    for part in INLINE.split(str(text)):
        if not part:
            continue
        if part.startswith("**"):
            yield part[2:-2], "b"
        elif part.startswith("`"):
            yield part[1:-1], "code"
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            yield part[1:-1], "i"
        else:
            yield part, ""


# ================================================================== DOCX
def _shade(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def _cell_margins(table, cm=0.12):
    tblPr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for side in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(int(cm * 567)))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tblPr.append(mar)


def _borders(table, color="CBD5E1", size=4):
    tblPr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:color"), color)
        b.append(el)
    tblPr.append(b)


def _repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def _add_runs(par, text, size=None, color=None, bold=None):
    for chunk, style in _runs(text):
        r = par.add_run(chunk)
        if style == "b":
            r.bold = True
        elif style == "i":
            r.italic = True
        elif style == "code":
            r.font.name = "Consolas"
            r.font.color.rgb = RGBColor.from_string("7C3AED")
        if size:
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = RGBColor.from_string(color)
        if bold:
            r.bold = True
    return par


def _field(par, instr):
    r = par.add_run()
    b = OxmlElement("w:fldChar"); b.set(qn("w:fldCharType"), "begin")
    t = OxmlElement("w:instrText"); t.set(qn("xml:space"), "preserve"); t.text = instr
    s = OxmlElement("w:fldChar"); s.set(qn("w:fldCharType"), "separate")
    txt = OxmlElement("w:t"); txt.text = "Right-click and choose Update Field to build the table of contents."
    e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end")
    for el in (b, t, s, txt, e):
        r._r.append(el)


def render_docx(doc: Doc, out: Path) -> Path:
    d = Document()
    d.core_properties.author = "FraudLens project team"
    d.core_properties.title = doc.title
    d.core_properties.subject = doc.subtitle
    sec = d.sections[0]
    if doc.landscape:
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
    else:
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Cm(2.0))
    sec.top_margin, sec.bottom_margin = Cm(1.8), Cm(1.8)
    usable = (sec.page_width - sec.left_margin - sec.right_margin) / 360000  # cm

    st = d.styles
    st["Normal"].font.name = "Calibri"
    st["Normal"].font.size = Pt(10.5)
    st["Normal"].paragraph_format.space_after = Pt(6)
    st["Normal"].paragraph_format.line_spacing = 1.12
    for name, size, color in (("Heading 1", 18, PALETTE["primary"]), ("Heading 2", 13.5, PALETTE["violet"]),
                              ("Heading 3", 11.5, PALETTE["teal"])):
        s = st[name]
        s.font.name = "Calibri"
        s.font.size = Pt(size)
        s.font.bold = True
        s.font.color.rgb = RGBColor.from_string(color)
        s.paragraph_format.space_before = Pt(14 if name == "Heading 1" else 10)
        s.paragraph_format.space_after = Pt(6)
        s.paragraph_format.keep_with_next = True

    # footer with page number
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _add_runs(fp, f"{doc.title}  |  page ", size=8, color=PALETTE["muted"])
    r = fp.add_run()
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar"); el.set(qn("w:fldCharType"), kind); r._r.append(el)
        else:
            el = OxmlElement("w:instrText"); el.text = text; r._r.append(el)
    r.font.size = Pt(8)

    # cover
    for _ in range(3):
        d.add_paragraph()
    band = d.add_table(rows=1, cols=1)
    band.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = band.cell(0, 0)
    _shade(c, PALETTE["primary"])
    _cell_margins(band, 0.6)
    p = c.paragraphs[0]
    _add_runs(p, doc.meta.get("kicker", "PROJECT REPORT"), size=10, color="C7D2FE", bold=True)
    p = c.add_paragraph()
    _add_runs(p, doc.title, size=26, color="FFFFFF", bold=True)
    if doc.subtitle:
        p = c.add_paragraph()
        _add_runs(p, doc.subtitle, size=13, color="E0E7FF")
    tags = [t for t in doc.meta.get("tags", "").split("|") if t]
    if tags:
        p = c.add_paragraph()
        _add_runs(p, "   ·   ".join(tags), size=9.5, color="C7D2FE", bold=True)
    d.add_paragraph()
    info = [(k, v) for k, v in doc.meta.items() if k not in ("kicker", "tags")]
    if info:
        t = d.add_table(rows=len(info), cols=2)
        _borders(t, "E2E8F0")
        _cell_margins(t, 0.18)
        for i, (k, v) in enumerate(info):
            _shade(t.cell(i, 0), PALETTE["tint"])
            _add_runs(t.cell(i, 0).paragraphs[0], k, bold=True, color=PALETTE["primary"])
            _add_runs(t.cell(i, 1).paragraphs[0], v)
            t.cell(i, 0).width, t.cell(i, 1).width = Cm(4.5), Cm(usable - 4.5)
    d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    fig_n = tab_n = 0
    for kind, val in doc.blocks:
        if kind in ("h1", "h2", "h3"):
            d.add_heading(val, level=int(kind[1]))
        elif kind == "p":
            _add_runs(d.add_paragraph(), val)
        elif kind in ("bullets", "numbers"):
            for item in val:
                _add_runs(d.add_paragraph(style="List Bullet" if kind == "bullets" else "List Number"), item)
        elif kind == "code":
            t = d.add_table(rows=1, cols=1)
            _shade(t.cell(0, 0), "0F172A")
            _cell_margins(t, 0.25)
            par = t.cell(0, 0).paragraphs[0]
            for k, line in enumerate(val.strip("\n").split("\n")):
                r = par.add_run(line)
                r.font.name, r.font.size = "Consolas", Pt(8.5)
                r.font.color.rgb = RGBColor.from_string("E2E8F0")
                if k < len(val.strip("\n").split("\n")) - 1:
                    r.add_break()
            d.add_paragraph()
        elif kind == "pagebreak":
            d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        elif kind == "toc":
            _add_runs(d.add_paragraph(), "Contents", size=18, color=PALETTE["primary"], bold=True)
            _field(d.add_paragraph(), 'TOC \\o "1-2" \\h \\z \\u')
            d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        elif kind == "callout":
            accent, bg, _ = CALLOUT[val["kind"]]
            t = d.add_table(rows=1, cols=1)
            _borders(t, accent, 8)
            _cell_margins(t, 0.25)
            cell = t.cell(0, 0)
            _shade(cell, bg)
            _add_runs(cell.paragraphs[0], val["title"], size=10.5, color=accent, bold=True)
            for para in str(val["text"]).split("\n"):
                _add_runs(cell.add_paragraph(), para, size=10)
            d.add_paragraph()
        elif kind == "kpis":
            n = len(val)
            t = d.add_table(rows=2, cols=n)
            _borders(t, "C7D2FE")
            _cell_margins(t, 0.15)
            for i, (value, label, *sub) in enumerate(val):
                top, bot = t.cell(0, i), t.cell(1, i)
                _shade(top, PALETTE["tint"]); _shade(bot, PALETTE["tint"])
                top.paragraphs[0].alignment = bot.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                _add_runs(top.paragraphs[0], value, size=17, color=PALETTE["violet"], bold=True)
                _add_runs(bot.paragraphs[0], label, size=8.5, color=PALETTE["ink"], bold=True)
                if sub and sub[0]:
                    q = bot.add_paragraph(); q.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    _add_runs(q, sub[0], size=7.5, color=PALETTE["muted"])
            d.add_paragraph()
        elif kind == "table":
            tab_n += 1
            cols, rows = val["columns"], val["rows"]
            if val["caption"]:
                cp = d.add_paragraph()
                _add_runs(cp, f"Table {tab_n}. {val['caption']}", size=9.5, color=PALETTE["primary"], bold=True)
                cp.paragraph_format.keep_with_next = True
            t = d.add_table(rows=1 + len(rows), cols=len(cols))
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            _borders(t)
            _cell_margins(t, 0.1)
            fsize = val["font"] or (9 if len(cols) <= 5 else 8)
            widths = val["widths"]
            if widths:
                scale = usable / sum(widths)
                widths = [w * scale for w in widths]
            for j, cname in enumerate(cols):
                cell = t.cell(0, j)
                _shade(cell, PALETTE["primary"])
                _add_runs(cell.paragraphs[0], cname, size=fsize, color="FFFFFF", bold=True)
            _repeat_header(t.rows[0])
            for i, row in enumerate(rows, start=1):
                hl = (i - 1) in val["highlight"]
                for j, v in enumerate(row):
                    cell = t.cell(i, j)
                    if hl:
                        _shade(cell, "EDE9FE")
                    elif i % 2 == 0:
                        _shade(cell, "F8FAFC")
                    _add_runs(cell.paragraphs[0], v, size=fsize,
                              bold=hl or (val["first_col_bold"] and j == 0))
            if widths:
                for row in t.rows:
                    for j, w in enumerate(widths):
                        row.cells[j].width = Cm(w)
            d.add_paragraph()
        elif kind == "figure":
            fig_n += 1
            if not val["path"].exists():
                _add_runs(d.add_paragraph(), f"[missing figure: {val['path'].name}]", color=PALETTE["red"])
                continue
            par = d.add_paragraph()
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
            par.paragraph_format.keep_with_next = True
            par.add_run().add_picture(str(val["path"]), width=Cm(min(val["width"], usable)))
            cp = d.add_paragraph()
            cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _add_runs(cp, f"Figure {fig_n}. {val['caption']}", size=9, color=PALETTE["muted"])
        elif kind == "flow":
            steps = val["steps"]
            t = d.add_table(rows=2, cols=len(steps))
            _borders(t, "FFFFFF", 12)
            _cell_margins(t, 0.12)
            colors = ["4F46E5", "6D28D9", "7C3AED", "0D9488", "0284C7", "16A34A", "D97706", "DB2777"]
            for j, (title, detail) in enumerate(steps):
                a, b = t.cell(0, j), t.cell(1, j)
                _shade(a, colors[j % len(colors)]); _shade(b, "F1F5F9")
                a.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                _add_runs(a.paragraphs[0], f"{j + 1}. {title}", size=8.5, color="FFFFFF", bold=True)
                _add_runs(b.paragraphs[0], detail, size=7.5)
            if val["caption"]:
                cp = d.add_paragraph(); cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                fig_n += 1
                _add_runs(cp, f"Figure {fig_n}. {val['caption']}", size=9, color=PALETTE["muted"])
            d.add_paragraph()
    out.parent.mkdir(parents=True, exist_ok=True)
    d.save(out)
    return out


# ================================================================== HTML
CSS = """
@page { size: A4 PORTRAIT; margin: 15mm 14mm 16mm 14mm;
  @bottom-center { content: "TITLE  ·  " counter(page); font: 8pt 'Segoe UI', sans-serif; color: #94A3B8; } }
@page cover { margin: 0; @bottom-center { content: none; } }
* { box-sizing: border-box; }
html { -webkit-print-color-adjust: exact; print-color-adjust: exact; }
body { font-family: 'Segoe UI', 'Inter', system-ui, sans-serif; color: #1E293B; font-size: 10pt; line-height: 1.5; margin: 0; }
.cover { page: cover; height: 297mm; width: 210mm; position: relative; overflow: hidden; color: #fff;
  background: radial-gradient(circle at 85% 12%, rgba(236,72,153,.55), transparent 38%),
              radial-gradient(circle at 10% 90%, rgba(13,148,136,.55), transparent 40%),
              linear-gradient(135deg, #1E1B4B 0%, #4338CA 45%, #7C3AED 100%);
  padding: 30mm 22mm; page-break-after: always; }
.cover .kicker { letter-spacing: .28em; font-size: 9pt; font-weight: 700; color: #C7D2FE; }
.cover h1 { font-size: 34pt; line-height: 1.1; margin: 8mm 0 5mm; font-weight: 800; }
.cover .sub { font-size: 14pt; color: #E0E7FF; max-width: 150mm; }
.cover .tags { margin-top: 10mm; }
.cover .tag { display: inline-block; padding: 2mm 4mm; border-radius: 99px; background: rgba(255,255,255,.14);
  border: 1px solid rgba(255,255,255,.3); margin: 0 2mm 2mm 0; font-size: 9pt; }
.cover .meta { position: absolute; left: 22mm; right: 22mm; bottom: 24mm; display: grid; grid-template-columns: 1fr 1fr; gap: 4mm 10mm; }
.cover .meta div { border-top: 1px solid rgba(255,255,255,.35); padding-top: 2mm; font-size: 9.5pt; }
.cover .meta b { display: block; font-size: 7.5pt; letter-spacing: .15em; color: #C7D2FE; text-transform: uppercase; }
.cover svg.mark { position: absolute; right: 18mm; top: 92mm; width: 70mm; opacity: .95; }
h1.sec { font-size: 19pt; color: #312E81; margin: 0 0 4mm; padding: 3mm 0 2mm; border-bottom: 3px solid #7C3AED;
  page-break-before: always; page-break-after: avoid; }
h1.sec .num { display: inline-block; background: linear-gradient(135deg,#4F46E5,#7C3AED); color: #fff; border-radius: 8px;
  padding: 0 3mm; margin-right: 3mm; font-size: 15pt; }
h2 { font-size: 13pt; color: #6D28D9; margin: 6mm 0 2mm; page-break-after: avoid; }
h3 { font-size: 11pt; color: #0F766E; margin: 4mm 0 1.5mm; page-break-after: avoid; }
p { margin: 0 0 2.6mm; text-align: justify; }
ul, ol { margin: 0 0 3mm 5mm; padding-left: 4mm; } li { margin-bottom: 1mm; }
code { font-family: Consolas, monospace; background: #F1F5F9; color: #6D28D9; padding: 0 1mm; border-radius: 3px; font-size: 9pt; }
pre { background: #0F172A; color: #E2E8F0; border-radius: 8px; padding: 3.5mm 4mm; font-size: 8.3pt; line-height: 1.45;
  white-space: pre-wrap; page-break-inside: avoid; }
.callout { border-radius: 10px; padding: 3mm 4mm 2mm; margin: 3mm 0 4mm; border-left: 5px solid; page-break-inside: avoid; }
.callout .t { font-weight: 700; font-size: 9.5pt; letter-spacing: .02em; margin-bottom: 1mm; }
.callout p { margin-bottom: 1.5mm; }
.kpis { display: grid; gap: 3mm; margin: 3mm 0 5mm; page-break-inside: avoid; }
.kpi { border-radius: 12px; padding: 3.5mm 3mm; color: #fff; text-align: center; }
.kpi .v { font-size: 18pt; font-weight: 800; line-height: 1.1; }
.kpi .l { font-size: 8.5pt; font-weight: 600; margin-top: 1mm; }
.kpi .s { font-size: 7.3pt; opacity: .85; }
table { width: 100%; border-collapse: separate; border-spacing: 0; margin: 1mm 0 5mm; font-size: 8.4pt;
  border: 1px solid #E2E8F0; border-radius: 8px; overflow: hidden; }
thead { display: table-header-group; }
th { background: linear-gradient(135deg,#4338CA,#6D28D9); color: #fff; text-align: left; padding: 1.8mm 2mm; font-weight: 600; }
td { padding: 1.6mm 2mm; border-top: 1px solid #EEF2F7; vertical-align: top; }
tr { page-break-inside: avoid; }
tbody tr:nth-child(even) td { background: #F8FAFC; }
tr.hl td { background: #EDE9FE !important; font-weight: 600; }
td.fc { font-weight: 600; color: #312E81; }
.cap { font-size: 8.6pt; color: #4338CA; font-weight: 700; margin: 3mm 0 1.2mm; page-break-after: avoid; }
figure { margin: 3mm 0 5mm; text-align: center; page-break-inside: avoid; }
figure img { max-width: 100%; border: 1px solid #E2E8F0; border-radius: 10px; padding: 2mm; background: #fff; }
figcaption { font-size: 8.6pt; color: #64748B; margin-top: 1.5mm; }
.flow { display: flex; gap: 0; margin: 3mm 0 2mm; page-break-inside: avoid; align-items: stretch; }
.flow .st { flex: 1; position: relative; margin-right: 3.2mm; border-radius: 10px; overflow: hidden; background: #F1F5F9; }
.flow .st:not(:last-child)::after { content: "▶"; position: absolute; right: -3.4mm; top: 3mm; color: #A5B4FC; font-size: 9pt; }
.flow .hd { color: #fff; font-weight: 700; font-size: 8.4pt; padding: 1.8mm 2mm; }
.flow .bd { font-size: 7.6pt; padding: 1.8mm 2mm; color: #334155; line-height: 1.35; }
.toc { columns: 2; column-gap: 10mm; font-size: 10pt; }
.toc div { break-inside: avoid; padding: 1mm 0; border-bottom: 1px dotted #CBD5E1; }
.toc .n { display: inline-block; min-width: 8mm; color: #7C3AED; font-weight: 700; }
.pb { page-break-after: always; }
"""
KPI_COLORS = ["#4F46E5", "#7C3AED", "#0D9488", "#DB2777", "#0284C7", "#D97706"]
FLOW_COLORS = ["#4F46E5", "#6D28D9", "#7C3AED", "#0D9488", "#0284C7", "#16A34A", "#D97706", "#DB2777"]
MARK = """<svg class="mark" viewBox="0 0 200 200" xmlns="http://www.w3.org/2000/svg">
<defs><linearGradient id="g" x1="0" x2="1" y1="0" y2="1"><stop offset="0" stop-color="#fff" stop-opacity=".95"/>
<stop offset="1" stop-color="#C4B5FD" stop-opacity=".7"/></linearGradient></defs>
<path d="M100 12 L172 40 V98 C172 140 140 172 100 188 C60 172 28 140 28 98 V40 Z" fill="none" stroke="url(#g)" stroke-width="7"/>
<circle cx="94" cy="92" r="34" fill="none" stroke="#fff" stroke-width="8"/>
<line x1="119" y1="117" x2="146" y2="144" stroke="#fff" stroke-width="10" stroke-linecap="round"/>
<path d="M76 94 l12 12 l24 -26" fill="none" stroke="#5EEAD4" stroke-width="8" stroke-linecap="round" stroke-linejoin="round"/>
</svg>"""


def _inline_html(text: str) -> str:
    out = []
    for chunk, style in _runs(text):
        e = html.escape(chunk)
        out.append({"b": f"<b>{e}</b>", "i": f"<i>{e}</i>", "code": f"<code>{e}</code>"}.get(style, e))
    return "".join(out)


def _img_src(path: Path) -> str:
    return "data:image/png;base64," + base64.b64encode(path.read_bytes()).decode()


def render_html(doc: Doc, out: Path) -> Path:
    css = CSS.replace("TITLE", doc.title.replace('"', "'"))
    if doc.landscape:
        css = css.replace("size: A4 PORTRAIT", "size: A4 landscape").replace("height: 297mm; width: 210mm", "height: 210mm; width: 297mm")
        css += ".cover svg.mark { top: 22mm; right: 24mm; width: 52mm; } .cover { padding: 22mm 24mm; } .cover .sub { max-width: 190mm; }"
    else:
        css = css.replace("size: A4 PORTRAIT", "size: A4")
        css += ".cover svg.mark { top: 138mm; right: 22mm; width: 64mm; }"  # below the title, tags and subtitle
    parts = [f"<!doctype html><html><head><meta charset='utf-8'><title>{html.escape(doc.title)}</title><style>{css}</style></head><body>"]
    tags = "".join(f"<span class='tag'>{html.escape(t)}</span>" for t in doc.meta.get("tags", "").split("|") if t)
    meta = "".join(f"<div><b>{html.escape(k)}</b>{_inline_html(v)}</div>" for k, v in doc.meta.items()
                   if k not in ("kicker", "tags"))
    parts.append(f"<section class='cover'><div class='kicker'>{html.escape(doc.meta.get('kicker', 'PROJECT REPORT'))}</div>"
                 f"<h1>{html.escape(doc.title)}</h1><div class='sub'>{html.escape(doc.subtitle)}</div>"
                 f"<div class='tags'>{tags}</div>{MARK}<div class='meta'>{meta}</div></section>")
    h1s = [v for k, v in doc.blocks if k == "h1"]
    fig_n = tab_n = sec_n = 0
    for kind, val in doc.blocks:
        if kind == "h1":
            sec_n += 1
            num = val.split(" ", 1)[0] if val[:1].isdigit() else ""
            title = val.split(" ", 1)[1] if num else val
            first = " style='page-break-before:auto'" if sec_n == 1 else ""
            parts.append(f"<h1 class='sec'{first}>" + (f"<span class='num'>{html.escape(num)}</span>" if num else "")
                         + f"{html.escape(title)}</h1>")
        elif kind in ("h2", "h3"):
            parts.append(f"<{kind}>{_inline_html(val)}</{kind}>")
        elif kind == "p":
            parts.append(f"<p>{_inline_html(val)}</p>")
        elif kind in ("bullets", "numbers"):
            tag = "ul" if kind == "bullets" else "ol"
            parts.append(f"<{tag}>" + "".join(f"<li>{_inline_html(i)}</li>" for i in val) + f"</{tag}>")
        elif kind == "code":
            parts.append(f"<pre>{html.escape(val.strip())}</pre>")
        elif kind == "pagebreak":
            parts.append("<div class='pb'></div>")
        elif kind == "toc":
            items = "".join(f"<div><span class='n'>{html.escape(h.split(' ', 1)[0]) if h[:1].isdigit() else '•'}</span>"
                            f"{html.escape(h.split(' ', 1)[1] if h[:1].isdigit() else h)}</div>" for h in h1s)
            parts.append(f"<h1 class='sec' style='page-break-before:auto'>Contents</h1><div class='toc'>{items}</div>")
        elif kind == "callout":
            accent, bg, _ = CALLOUT[val["kind"]]
            body = "".join(f"<p>{_inline_html(x)}</p>" for x in str(val["text"]).split("\n"))
            parts.append(f"<div class='callout' style='border-color:#{accent};background:#{bg}'>"
                         f"<div class='t' style='color:#{accent}'>{html.escape(val['title'])}</div>{body}</div>")
        elif kind == "kpis":
            cells = "".join(
                f"<div class='kpi' style='background:linear-gradient(135deg,{KPI_COLORS[i % 6]},{KPI_COLORS[(i + 1) % 6]})'>"
                f"<div class='v'>{html.escape(str(v[0]))}</div><div class='l'>{html.escape(v[1])}</div>"
                + (f"<div class='s'>{html.escape(v[2])}</div>" if len(v) > 2 and v[2] else "") + "</div>"
                for i, v in enumerate(val))
            parts.append(f"<div class='kpis' style='grid-template-columns:repeat({len(val)},1fr)'>{cells}</div>")
        elif kind == "table":
            tab_n += 1
            cap = f"<div class='cap'>Table {tab_n}. {_inline_html(val['caption'])}</div>" if val["caption"] else ""
            ws = val["widths"]
            colgroup = ""
            if ws:
                tot = sum(ws)
                colgroup = "<colgroup>" + "".join(f"<col style='width:{w / tot * 100:.1f}%'>" for w in ws) + "</colgroup>"
            head = "".join(f"<th>{_inline_html(c)}</th>" for c in val["columns"])
            body = []
            for i, row in enumerate(val["rows"]):
                cls = " class='hl'" if i in val["highlight"] else ""
                tds = "".join(f"<td{' class=fc' if (val['first_col_bold'] and j == 0) else ''}>{_inline_html(c)}</td>"
                              for j, c in enumerate(row))
                body.append(f"<tr{cls}>{tds}</tr>")
            fs = f" style='font-size:{val['font']}pt'" if val["font"] else ""
            parts.append(f"{cap}<table{fs}>{colgroup}<thead><tr>{head}</tr></thead><tbody>{''.join(body)}</tbody></table>")
        elif kind == "figure":
            fig_n += 1
            if val["path"].exists():
                w = min(100, val["width"] / 17.0 * 100)
                parts.append(f"<figure><img src='{_img_src(val['path'])}' style='width:{w:.0f}%'>"
                             f"<figcaption><b>Figure {fig_n}.</b> {_inline_html(val['caption'])}</figcaption></figure>")
        elif kind == "flow":
            steps = "".join(f"<div class='st'><div class='hd' style='background:{FLOW_COLORS[j % 8]}'>{j + 1}. "
                            f"{html.escape(t)}</div><div class='bd'>{_inline_html(dsc)}</div></div>"
                            for j, (t, dsc) in enumerate(val["steps"]))
            parts.append(f"<div class='flow'>{steps}</div>")
            if val["caption"]:
                fig_n += 1
                parts.append(f"<figure style='margin-top:0'><figcaption><b>Figure {fig_n}.</b> "
                             f"{_inline_html(val['caption'])}</figcaption></figure>")
    parts.append("</body></html>")
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text("".join(parts), encoding="utf-8")
    return out


def html_to_pdf(html_path: Path, pdf_path: Path) -> Path:
    """Print the HTML with headless Microsoft Edge / Chrome (both render print CSS faithfully)."""
    candidates = [r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe",
                  r"C:\Program Files\Microsoft\Edge\Application\msedge.exe",
                  r"C:\Program Files\Google\Chrome\Application\chrome.exe"]
    exe = next((c for c in candidates if Path(c).exists()), None)
    if exe is None:
        raise RuntimeError("Edge or Chrome is required to print the PDF")
    subprocess.run([exe, "--headless=new", "--disable-gpu", "--no-pdf-header-footer", "--virtual-time-budget=8000",
                    f"--print-to-pdf={pdf_path}", html_path.resolve().as_uri()], check=True, timeout=240,
                   capture_output=True)
    return pdf_path


def update_docx_fields_and_export(docx_path: Path, pdf_path: Path | None = None) -> None:
    """Use Microsoft Word (if installed) to refresh the table of contents; optionally export a PDF."""
    ps = (f"$w = New-Object -ComObject Word.Application; $w.Visible = $false; "
          f"$d = $w.Documents.Open('{docx_path.resolve()}'); "
          f"foreach ($t in $d.TablesOfContents) {{ $t.Update() }}; $d.Save(); "
          + (f"$d.ExportAsFixedFormat('{pdf_path.resolve()}', 17); " if pdf_path else "")
          + "$d.Close(); $w.Quit()")
    subprocess.run(["powershell", "-NoProfile", "-Command", ps], check=False, timeout=300, capture_output=True)
